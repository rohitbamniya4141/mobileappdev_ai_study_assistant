"""Tests for DOCX support in document_service.py."""
import io

import pytest
from docx import Document as DocxDocument

from app.services.document_service import (
    InvalidDocumentError,
    extract_text_from_docx,
    extract_text_from_pdf,
    validate_docx,
)


def _make_docx(paragraphs, table_rows=None):
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=max(len(r) for r in table_rows))
        for i, row in enumerate(table_rows):
            for j, cell_text in enumerate(row):
                table.rows[i].cells[j].text = cell_text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_validate_docx_accepts_real_docx():
    validate_docx(_make_docx(["Hello world"]))

def test_validate_docx_rejects_random_bytes():
    with pytest.raises(InvalidDocumentError, match="not a valid DOCX"):
        validate_docx(b"this is not a docx file at all")

def test_validate_docx_rejects_pdf_magic():
    with pytest.raises(InvalidDocumentError, match="not a valid DOCX"):
        validate_docx(b"%PDF-1.4 fake content")

def test_validate_docx_rejects_empty():
    with pytest.raises(InvalidDocumentError):
        validate_docx(b"")

def test_extract_text_from_docx_paragraphs():
    text = extract_text_from_docx(_make_docx(["Para one.", "Para two."]))
    assert "Para one." in text
    assert "Para two." in text

def test_extract_text_from_docx_table():
    rows = [["Name", "Age"], ["Alice", "30"]]
    text = extract_text_from_docx(_make_docx(["Header"], table_rows=rows))
    assert "Name" in text and "Alice" in text and "30" in text

def test_extract_text_from_docx_table_pipe_separator():
    text = extract_text_from_docx(_make_docx([], table_rows=[["Python", "FastAPI"]]))
    assert "|" in text

def test_extract_text_from_docx_empty_raises():
    with pytest.raises(InvalidDocumentError, match="No extractable text"):
        extract_text_from_docx(_make_docx([]))

def test_extract_text_from_docx_rejects_non_docx():
    with pytest.raises(InvalidDocumentError, match="not a valid DOCX"):
        extract_text_from_docx(b"fake bytes")

def test_pdf_still_rejects_docx_bytes():
    with pytest.raises(InvalidDocumentError, match="not a valid PDF"):
        extract_text_from_pdf(_make_docx(["Some text"]))
